"""Socket-free checks for the local app's security and lawyer workflow."""

import time
from dataclasses import replace
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient

from nopoach.app.main import create_app, launch
from nopoach.market.define import Geography, Market
from nopoach.store.workspace import Workspace


@pytest.fixture
def client():
    app = create_app(port=8765, token="synthetic-session")
    with TestClient(app, base_url="http://127.0.0.1:8765") as browser:
        browser.get("/?token=synthetic-session")
        yield browser


def wait_screen(client):
    for _ in range(500):
        if client.app.state.session.status != "Collecting":
            return client.get("/screen")
        time.sleep(0.01)
    pytest.fail("Background screen did not finish")


def test_security_and_one_time_token():
    app = create_app(port=8765, token="synthetic-session")
    with TestClient(app, base_url="http://127.0.0.1:8765") as browser:
        assert browser.get("/").status_code == 403
        response = browser.get("/?token=synthetic-session", follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"] == "/"
        assert "HttpOnly" in response.headers["set-cookie"]
        assert "SameSite=strict" in response.headers["set-cookie"]
        assert browser.get("/").status_code == 200
        assert (
            browser.get("/", headers={"host": "foreign.invalid:8765"}).status_code
            == 403
        )
        assert browser.get("/", headers={"host": "127.0.0.1:9999"}).status_code == 403
        assert (
            browser.post(
                "/workspace", headers={"origin": "https://foreign.invalid"}
            ).status_code
            == 403
        )
        assert browser.post("/workspace").status_code == 403
        assert browser.post("/workspace", headers={"origin": "null"}).status_code == 403
        browser.cookies.clear()
        assert browser.get("/?token=synthetic-session").status_code == 403


def test_launcher_only_binds_loopback():
    from threading import Event

    browser_opened = Event()
    with (
        patch("nopoach.app.main.uvicorn.run"),
        patch("nopoach.app.main.uvicorn.Server") as server_type,
        patch("nopoach.app.main.webbrowser.open") as opened,
    ):
        server = server_type.return_value
        server.started = False
        opened.side_effect = lambda url: browser_opened.set()

        def serve():
            assert not opened.called  # Do not open a browser before the listener.
            server.started = True
            assert browser_opened.wait(2)

        server.run.side_effect = serve
        launch(port=8765)
        server_type.assert_called_once()
        configuration = server_type.call_args.args[0]
        assert configuration.host == "127.0.0.1"
        assert configuration.port == 8765
        assert configuration.access_log is False
        assert opened.call_args.args[0].startswith("http://127.0.0.1:8765/?token=")


def test_pages_assets_and_csp(client):
    response = client.get("/")
    assert "default-src 'self'" in response.headers["content-security-policy"]
    assert response.headers["referrer-policy"] == "same-origin"
    assert "Skip to content" in response.text
    assert client.get("/static/app.css").status_code == 200
    assert client.get("/static/fonts/SourceSerif4-Variable.woff2").status_code == 200
    assert (
        client.get("/market").status_code == 303
        or "Choose a workspace" in client.get("/market").text
    )


def open_workspace(client, tmp_path):
    root = tmp_path / "synthetic-workspace"
    with Workspace.open(root) as workspace:
        from nopoach.imports.loader import import_records

        import_records(
            workspace,
            [
                {
                    "person_key": "synthetic-director",
                    "employer_name": name,
                    "title": "Director",
                    "location": "Fixture",
                    "start_date": "2015",
                    "start_precision": "year",
                    "end_date": "2024",
                    "end_precision": "year",
                }
                for name in ("Synthetic A", "Synthetic B")
            ],
            source="Synthetic source",
            provenance_class="synthetic",
        )
    response = client.post(
        "/workspace",
        data={"folder": str(root)},
        headers={"origin": "http://127.0.0.1:8765"},
    )
    assert response.status_code == 200
    return root


def run_market(client):
    return client.post(
        "/market",
        data={
            "occupations": "29-1141",
            "geography_kind": "metro",
            "geography_code": "SYNTHETIC",
            "start_year": "2015",
            "end_year": "2024",
            "employers": [r[0] for r in client.app.state.session.employers],
        },
        headers={"origin": "http://127.0.0.1:8765"},
    )


def test_import_rejects_rerun_partial_and_dossier(client, tmp_path):
    import io

    from docx import Document
    from lxml import html

    root = open_workspace(client, tmp_path)
    run_market(client)
    response = wait_screen(client)
    assert "Unreliable pairs" in response.text
    assert "insufficient" in response.text
    assert "government or administrative records" in response.text
    # U10 keeps independent ties visible; U12 must still isolate unreliable pairs.
    assert (
        "Synthetic A / Synthetic B"
        not in response.text.split('<details id="unreliable">')[0]
    )
    confirmed = client.app.state.session.market.employer_ids
    assert client.get("/screen/accept").status_code == 200
    response = client.post(
        "/screen/accept",
        data={"confirm": "yes"},
        headers={"origin": "http://127.0.0.1:8765"},
    )
    assert "Partial screen" in wait_screen(client).text
    pair = client.app.state.session.result.unreliable[0].pair
    dossier = client.get(f"/pairs/{pair.employer_a_id}/{pair.employer_b_id}")
    assert "Explanations checked" in dossier.text
    assert "Synthetic source" in dossier.text
    export_links = html.fromstring(dossier.text).xpath(
        "//a[normalize-space(.)='Export dossier (Word)']/@href"
    )
    assert export_links == [f"/pairs/{pair.employer_a_id}/{pair.employer_b_id}/export"]
    exported = client.get(export_links[0])
    assert exported.status_code == 200
    assert exported.headers["content-disposition"] == (
        'attachment; filename="dossier.docx"'
    )
    files = list((root / "exports").glob("*.docx"))
    assert len(files) == 1
    assert files[0].read_bytes() == exported.content
    document = Document(io.BytesIO(exported.content))
    assert "Partial screen" in "\n".join(p.text for p in document.paragraphs)
    assert "Partial screen" in dossier.text
    body = (
        "worker,employer_name,title,location,start_date,start_precision,end_date,end_precision\n"
        "synthetic-new,Synthetic A,Registered nurse,Fixture,2015,year,2024,year\n"
        "synthetic-reject,Synthetic A,Registered nurse,Fixture,2024,year,2010,year\n"
    )
    response = client.post(
        "/imports/preview",
        files={"file": ("synthetic.csv", body)},
        headers={"origin": "http://127.0.0.1:8765"},
    )
    assert "Map columns" in response.text
    response = client.post(
        "/imports",
        data={
            "source": "Synthetic supplemental",
            "provenance_class": "government_or_administrative_records",
            "map_person_key": "worker",
        },
        headers={"origin": "http://127.0.0.1:8765"},
    )
    assert "end_date before start_date" in response.text
    assert "Rejected row 2" in response.text
    wait_screen(client)
    assert client.app.state.session.market.employer_ids == confirmed
    assert "Synthetic supplemental" in client.get("/imports").text


def test_dossier_export_unavailable_when_public_function_is_absent(client, tmp_path):
    from types import SimpleNamespace

    from lxml import html

    root = open_workspace(client, tmp_path)
    run_market(client)
    wait_screen(client)
    pair = client.app.state.session.result.unreliable[0].pair
    url = f"/pairs/{pair.employer_a_id}/{pair.employer_b_id}"
    with patch(
        "nopoach.app.routes.pages.importlib.import_module",
        return_value=SimpleNamespace(),
    ):
        dossier = client.get(url)
        assert "Export unavailable." in dossier.text
        assert not html.fromstring(dossier.text).xpath(
            "//a[normalize-space(.)='Export dossier (Word)']"
        )
        exported = client.get(url + "/export")
    assert exported.status_code == 503
    assert "Export unavailable" in exported.text
    assert not (root / "exports").exists()


def test_failed_collector_keeps_results_and_can_retry(tmp_path):
    attempts = []

    def collector(workspace, market):
        attempts.append(1)
        if len(attempts) == 1:
            raise RuntimeError("Synthetic index did not download")

    app = create_app(
        port=8765, token="synthetic-session", collectors={"Synthetic index": collector}
    )
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        client.get("/?token=synthetic-session")
        open_workspace(client, tmp_path)
        run_market(client)
        response = wait_screen(client)
        assert "Synthetic index did not download" in response.text
        assert "Retry source" in response.text
        assert app.state.session.status == "Failed"
        assert app.state.session.result.unreliable
        client.post(
            "/screen/retry",
            data={"source": "Synthetic index"},
            headers={"origin": "http://127.0.0.1:8765"},
        )
        wait_screen(client)
        assert attempts == [1, 1]
        assert app.state.session.status == "Insufficient"


@pytest.mark.network_off
def test_no_network_for_stored_screen(client, tmp_path):
    open_workspace(client, tmp_path)
    with patch(
        "socket.socket.connect", side_effect=AssertionError("Network forbidden")
    ):
        run_market(client)
        assert wait_screen(client).status_code == 200


def test_one_way_tab_excludes_mutual(client, tmp_path):
    open_workspace(client, tmp_path)
    run_market(client)
    wait_screen(client)
    result = client.app.state.session.result
    pair = result.unreliable[0].pair
    from nopoach.rank import wording

    one_way = replace(pair, label=wording.ONE_WAY_LABEL)
    client.app.state.session.result = replace(
        result, mutual=(), one_way=(one_way,), unreliable=()
    )
    assert "Synthetic A / Synthetic B" not in client.get("/screen").text
    assert wording.ONE_WAY_LABEL in client.get("/screen?tab=one-way").text


def test_market_validation(client, tmp_path):
    open_workspace(client, tmp_path)
    response = client.post(
        "/market",
        data={
            "occupations": "bad",
            "geography_kind": "metro",
            "geography_code": "SYNTHETIC",
            "start_year": "2015",
            "end_year": "2024",
        },
        headers={"origin": "http://127.0.0.1:8765"},
    )
    assert response.status_code == 400
    assert "Occupation must use SOC" in response.text
    assert client.app.state.session.market is None


def test_real_pipeline_empty_workspace(client, tmp_path):
    root = tmp_path / "empty"
    client.post(
        "/workspace",
        data={"folder": str(root)},
        headers={"origin": "http://127.0.0.1:8765"},
    )
    client.app.state.session.market = Market(
        frozenset({"29-1141"}), Geography("metro", "SYNTHETIC"), 2015, 2024
    )
    client.post("/screen/run", headers={"origin": "http://127.0.0.1:8765"})
    assert "No pairs met the flagging threshold" in wait_screen(client).text


def test_nonloopback_peer_refused_without_a_socket():
    app = create_app(port=8765, token="synthetic-session")
    with TestClient(
        app, base_url="http://127.0.0.1:8765", client=("192.0.2.12", 10001)
    ) as browser:
        response = browser.get("/?token=synthetic-session")
        assert response.status_code == 403
        assert "Loopback requests only" in response.text


@pytest.mark.parametrize(
    ("method", "headers", "peer", "reason"),
    [
        ("GET", {}, "testclient", "Session token required"),
        ("GET", {"host": "foreign.invalid:8765"}, "testclient", "Foreign Host refused"),
        ("GET", {}, "192.0.2.12", "Loopback requests only"),
        ("POST", {"origin": "null"}, "testclient", "Same-origin requests only"),
        (
            "POST",
            {"origin": "https://foreign.invalid"},
            "testclient",
            "Same-origin requests only",
        ),
        ("POST", {}, "testclient", "Same-origin requests only"),
    ],
)
def test_refusals_are_accessible_html_without_screen_data(
    method, headers, peer, reason
):
    from lxml import html

    app = create_app(port=8765, token="synthetic-session")
    # Refusals must not disclose the current market to an unauthenticated request.
    app.state.session.market = Market(
        frozenset({"29-1141"}), Geography("metro", "PRIVATE-MARKET"), 2015, 2024
    )
    with TestClient(
        app, base_url="http://127.0.0.1:8765", client=(peer, 10001)
    ) as browser:
        response = browser.request(method, "/", headers=headers)
    assert response.status_code == 403
    assert response.headers["content-type"].startswith("text/html")
    assert response.headers["referrer-policy"] == "same-origin"
    assert "default-src 'self'" in response.headers["content-security-policy"]
    tree = html.fromstring(response.text)
    assert tree.get("lang") == "en"
    assert tree.xpath("//title/text()") == ["Request refused · Labor market screen"]
    assert tree.xpath("//main[@id='content']//h1/text()") == ["Request refused"]
    assert reason in tree.xpath("//main")[0].text_content()
    assert "PRIVATE-MARKET" not in response.text
    assert "synthetic-session" not in response.text


@pytest.mark.parametrize(
    ("scenario", "occupation"),
    [("AE1", "29-1141"), ("AE2", "29-1141"), ("AE3", "51-3022")],
)
def test_real_acceptance_markets_through_routes(client, tmp_path, scenario, occupation):
    import runpy
    from pathlib import Path

    from nopoach.synth.scenarios import named_scenario

    helper = runpy.run_path(str(Path(__file__).parents[1] / "test_pipeline.py"))[
        "load_resolved_scenario"
    ]
    root = tmp_path / scenario
    with Workspace.open(root) as workspace:
        helper(workspace, named_scenario(scenario))
    headers = {"origin": "http://127.0.0.1:8765"}
    client.post("/workspace", data={"folder": str(root)}, headers=headers)
    client.post(
        "/market",
        data={
            "occupations": occupation,
            "geography_kind": "metro",
            "geography_code": "SYNTHETIC",
            "start_year": "2015",
            "end_year": "2024",
            "action": "propose",
        },
        headers=headers,
    )
    # Proposal uses U6 estimates rather than guessing competitors from name.
    assert client.app.state.session.market.employer_ids is None
    client.post(
        "/market",
        data={
            "occupations": occupation,
            "geography_kind": "metro",
            "geography_code": "SYNTHETIC",
            "start_year": "2015",
            "end_year": "2024",
            "employers": [r[0] for r in client.app.state.session.employers],
        },
        headers=headers,
    )
    response = wait_screen(client)
    from lxml import html

    tree = html.fromstring(response.text)
    reliable_rows = tree.xpath("(//table)[1]/tbody/tr")
    if scenario == "AE1":
        assert (
            "Synthetic Amber Hospital / Synthetic Birch Hospital"
            in reliable_rows[0].text_content()
        )
        result = client.app.state.session.result
        pair = result.mutual[0]
        # Presentation must stay chronological even when stored evidence arrives
        # in a different order. The direction indexes still identify the same flow.
        pair = replace(
            pair,
            flow=replace(
                pair.flow,
                directions=tuple(
                    replace(direction, years=tuple(reversed(direction.years)))
                    for direction in pair.flow.directions
                ),
            ),
        )
        client.app.state.session.result = replace(
            result, mutual=(pair, *result.mutual[1:])
        )
        url = f"/pairs/{pair.employer_a_id}/{pair.employer_b_id}"
        dossier = client.get(url)
        assert "Hollow: expected. Solid: observed." in dossier.text
        assert "Break interval: 2019–2020" in dossier.text
        dossier_tree = html.fromstring(dossier.text)
        for direction_index in (0, 1):
            flow_data = dossier_tree.xpath("//section[@id='flow']//details")[
                direction_index
            ]
            assert flow_data.xpath(".//tbody/tr/th/text()") == [
                str(year) for year in range(2015, 2025)
            ]
            links = flow_data.xpath(".//a[@hx-get]")
            assert links[0].get("hx-get") == (
                url + f"/moves?direction={direction_index}&year=2015"
            )
            assert links[0].get("hx-target") == f"#moves-{direction_index}-2015"
            assert flow_data.get("id") == f"flow-data-{direction_index}"
        moves = client.get(
            url + "/moves?direction=0&year=2015", headers={"HX-Request": "true"}
        )
        assert "Person-level moves in 2015" in moves.text
        assert "Person key" in moves.text
        assert "Position records" in moves.text
        assert "<html" not in moves.text
    elif scenario == "AE2":
        assert "Synthetic Amber Hospital / Synthetic Birch Hospital" not in " ".join(
            row.text_content() for row in reliable_rows
        )
        one_way = client.get("/screen?tab=one-way")
        assert "Synthetic Amber Hospital / Synthetic Birch Hospital" in one_way.text
        assert "Possible unilateral no-hire policy" in one_way.text
    else:
        assert "Missing worker flow: <span" in response.text
        assert "Unreliable pairs (1)" in response.text
        assert "below 20%" in response.text
        sources = tree.xpath("//div[@class='source']/h4/text()")
        assert sources[0] == "State unemployment-insurance wage records"


def test_workforce_coverage_has_named_keyboard_scroll_region(client, tmp_path):
    from lxml import html

    open_workspace(client, tmp_path)
    run_market(client)
    response = wait_screen(client)
    tree = html.fromstring(response.text)
    coverage = tree.xpath(
        "//div[@role='region' and "
        "@aria-label='Workforce coverage by employer and year']"
    )
    assert len(coverage) == 1
    assert coverage[0].get("tabindex") == "0"
    assert coverage[0].xpath("./table/thead/tr/th[last()]/text()") == [
        "Denominator source"
    ]


def test_parent_override_with_reason_keeps_confirmed_employers(client, tmp_path):
    root = open_workspace(client, tmp_path)
    run_market(client)
    wait_screen(client)
    session = client.app.state.session
    confirmed = session.market.employer_ids
    a, b = (r[0] for r in session.employers)
    client.post(
        "/imports/parent",
        data={
            "employer_id": a,
            "parent_id": b,
            "valid_from": "2020-01-01",
            "reason": "Synthetic ownership evidence",
        },
        headers={"origin": "http://127.0.0.1:8765"},
    )
    wait_screen(client)
    assert session.market.employer_ids == confirmed
    with Workspace.open(root) as workspace:
        from datetime import date

        from nopoach.resolve.parents import ParentResolver

        evidence = ParentResolver(workspace).evidence(a, date(2021, 1, 1))
        assert evidence[0]["parent_id"] == b
        assert evidence[0]["origin"] == "manual_override"
        assert evidence[0]["reason"] == "Synthetic ownership evidence"


def test_export_uses_public_function_with_accepted_partial_result(client, tmp_path):
    from types import SimpleNamespace

    open_workspace(client, tmp_path)
    run_market(client)
    wait_screen(client)
    client.post(
        "/screen/accept",
        data={"confirm": "yes"},
        headers={"origin": "http://127.0.0.1:8765"},
    )
    wait_screen(client)
    pair = client.app.state.session.result.unreliable[0].pair
    captured = []

    def export_dossier(*, workspace, screen, pair, destination):
        captured.append((workspace.root, screen, pair))
        destination.write_bytes(b"synthetic export contract fixture")

    with patch(
        "nopoach.app.routes.pages.importlib.import_module",
        return_value=SimpleNamespace(export_dossier=export_dossier),
    ):
        response = client.get(
            f"/pairs/{pair.employer_a_id}/{pair.employer_b_id}/export"
        )
    assert response.status_code == 200
    assert captured[0][1].accept_partial is True
    assert captured[0][1].partial_screen is True
    assert captured[0][2] == pair
    assert (
        response.headers["content-disposition"] == 'attachment; filename="dossier.docx"'
    )


def test_filters_score_sort_headers_and_paging(client, tmp_path):
    from lxml import html

    open_workspace(client, tmp_path)
    run_market(client)
    wait_screen(client)
    result = client.app.state.session.result
    example = result.unreliable[0].pair
    leads = tuple(
        replace(
            example,
            employer_a_id=f"synthetic-{i}",
            employer_a_name=f"Synthetic Employer {i:02}",
            score=i / 100,
        )
        for i in range(1, 26)
    )
    client.app.state.session.result = replace(
        result, mutual=leads, one_way=(), unreliable=()
    )
    tree = html.fromstring(client.get("/screen").text)
    rows = tree.xpath("(//table)[1]/tbody/tr")
    assert "Synthetic Employer 25" in rows[0].text_content()
    assert len(rows) == 20
    assert "Page 1 of 2" in tree.text_content()
    second = client.get("/screen?page_number=2")
    assert "Synthetic Employer 05" in second.text
    assert "Synthetic Employer 25" not in second.text
    named = client.get("/screen?sort=employer")
    assert 'aria-sort="ascending"' in named.text
    assert (
        "Synthetic Employer 01"
        in html.fromstring(named.text).xpath("(//table)[1]/tbody/tr")[0].text_content()
    )
    filtered = client.get("/screen?employer=Employer%2001&signal=ties")
    assert "Synthetic Employer 01" in filtered.text
    assert "Synthetic Employer 02" not in filtered.text
    assert (
        "No pairs met the flagging threshold" in client.get("/screen?signal=flow").text
    )


def test_import_requires_source_and_provenance(client, tmp_path):
    open_workspace(client, tmp_path)
    body = (
        "person_key,employer_name,title,location,start_date,start_precision,end_date,end_precision\n"
        "synthetic-new,Synthetic A,Registered nurse,Fixture,2015,year,2024,year\n"
    )
    headers = {"origin": "http://127.0.0.1:8765"}
    client.post(
        "/imports/preview", files={"file": ("synthetic.csv", body)}, headers=headers
    )
    missing_class = client.post(
        "/imports", data={"source": "Synthetic import"}, headers=headers
    )
    assert missing_class.status_code == 400
    assert "Choose a provenance class" in missing_class.text
    missing_source = client.post(
        "/imports", data={"provenance_class": "synthetic"}, headers=headers
    )
    assert missing_source.status_code == 400
    assert "A source name is required" in missing_source.text


def test_clause_dossier_includes_filing_and_text_record(client, tmp_path):
    from datetime import date

    from nopoach.market.define import Employer
    from nopoach.signals.clauses import ClauseDocument, store_document

    root = open_workspace(client, tmp_path)
    with Workspace.open(root) as workspace:
        employers = tuple(Employer(*r) for r in client.app.state.session.employers)
        receipt = workspace.provenance.append(
            source="Synthetic filed contract",
            url="file:synthetic-contract",
            method="IMPORT",
            body=b"synthetic",
            collection_method="synthetic",
        )
        store_document(
            workspace,
            ClauseDocument(
                "Synthetic A and Synthetic B agree that neither party shall hire "
                "the other party's employees during the term of this agreement.",
                receipt,
                date(2020, 1, 1),
                "synthetic_contract",
                employers[0],
                document_name="Synthetic exhibit",
            ),
            employers,
        )
    run_market(client)
    wait_screen(client)
    pair = client.app.state.session.result.unreliable[0].pair
    dossier = client.get(f"/pairs/{pair.employer_a_id}/{pair.employer_b_id}")
    assert "neither party shall hire" in dossier.text
    assert "Synthetic exhibit (characters" in dossier.text
    assert "Synthetic filed contract" in dossier.text
    assert "binding in the text: mutual" in dossier.text
    assert "Combined score calculation" in dossier.text
    assert dossier.status_code == 200


def test_changed_market_does_not_display_previous_market_results(client, tmp_path):
    open_workspace(client, tmp_path)
    run_market(client)
    wait_screen(client)
    headers = {"origin": "http://127.0.0.1:8765"}
    with patch(
        "nopoach.app.routes.pages.run_screen",
        side_effect=RuntimeError("Synthetic scoring failure"),
    ):
        client.post(
            "/market",
            data={
                "occupations": "15-1252",
                "geography_kind": "metro",
                "geography_code": "SYNTHETIC-NEW",
                "start_year": "2015",
                "end_year": "2024",
                "employers": [r[0] for r in client.app.state.session.employers],
            },
            headers=headers,
        )
        response = wait_screen(client)
    assert client.app.state.session.result is None
    assert "Synthetic scoring failure" in response.text
    assert "Synthetic A / Synthetic B" not in response.text
