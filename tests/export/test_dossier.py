import json
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pyarrow as pa
import pytest
from docx import Document
from matplotlib.figure import Figure

from nopoach.export.dossier import export_dossier
from nopoach.imports.loader import import_records
from nopoach.market.define import Employer, Geography, Market
from nopoach.pipeline import run_screen
from nopoach.signals.clauses import ClauseDocument, store_document
from nopoach.store.workspace import Workspace
from nopoach.synth.generator import generate_market
from nopoach.synth.scenarios import named_scenario


@pytest.fixture(scope="module")
def ae1(tmp_path_factory):
    root = tmp_path_factory.mktemp("dossier-ae1")
    with Workspace.open(root / "workspace") as workspace:
        scenario = named_scenario("AE1")
        generated = generate_market(scenario, seed=42)
        receipt = workspace.provenance.append(
            source="synthetic",
            url="file:AE1",
            method="IMPORT",
            body=json.dumps(generated.rows, sort_keys=True).encode(),
            collection_method="synthetic",
        )
        db = workspace.connection
        db.register("fixture", pa.Table.from_pylist(list(generated.rows)))
        db.register("truth", pa.Table.from_pylist(list(generated.true_rows)))
        db.execute(
            "INSERT INTO employers SELECT DISTINCT employer_id, employer_name, "
            "NULL, ? FROM fixture",
            [receipt.id],
        )
        db.execute(
            "INSERT INTO people SELECT DISTINCT person_key, person_name, ? "
            "FROM fixture",
            [receipt.id],
        )
        db.execute(
            "CREATE TEMP TABLE jobs AS SELECT 'job-' || row_number() "
            "OVER () AS id, * FROM fixture"
        )
        db.execute(
            "INSERT INTO positions SELECT id, person_key, employer_id, title, "
            "occupation_code, start_date::DATE, nullif(end_date, '')::DATE, "
            "start_precision, end_precision, ? FROM jobs",
            [receipt.id],
        )
        db.execute(
            "INSERT INTO position_worksites SELECT id, cbsa, latitude::DOUBLE, "
            "longitude::DOUBLE, 'synthetic', 'supplied_cbsa_centroid', ? FROM jobs",
            [receipt.id],
        )
        for year in range(2015, 2025):
            db.execute(
                "INSERT INTO workforce_inputs SELECT 'count-' || employer_id || ?, "
                "employer_id, ?, '[\"29-1141\"]', 'metro', 'SYNTHETIC', 'imported', "
                "count(*), count(*), count(*), 'synthetic', ?::DATE, ? "
                "FROM truth WHERE start_date::DATE <= ?::DATE AND "
                "(end_date = '' OR end_date::DATE >= ?::DATE) GROUP BY employer_id",
                [
                    str(year),
                    year,
                    f"{year}-01-01",
                    receipt.id,
                    f"{year}-01-01",
                    f"{year}-01-01",
                ],
            )
        db.execute("DROP TABLE jobs")
        db.unregister("fixture")
        db.unregister("truth")
        market = Market(
            frozenset({"29-1141"}), Geography("metro", "SYNTHETIC"), 2015, 2024
        )
        screen = run_screen(workspace, market, accept_partial=True)
        path = export_dossier(workspace, screen, screen.mutual[0], root / "AE1.docx")
        yield workspace, market, screen, path


def text_of(document):
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )


def test_ae1_standalone_editable_document_has_all_sections(ae1):
    _, _, screen, path = ae1
    assert path.is_file()
    document = Document(path)
    assert document.paragraphs[0].style.name == "Title"
    headings = {p.text for p in document.paragraphs if p.style.name == "Heading 1"}
    assert headings == {
        "Plain language summary",
        "Worker flow",
        "Filed clause excerpts",
        "Executive and board ties",
        "Explanation checks",
        "Coverage and sufficiency",
        "Data sources",
        "Method note",
    }
    text = text_of(document)
    assert "Synthetic Amber Hospital" in text and "Synthetic Birch Hospital" in text
    assert "warrants further inquiry" in text
    assert "Geographic separation" in text and "Non-compete prevalence" in text
    assert "not probabilities" in text
    assert "2015" in text and "2024" in text and "29-1141" in text
    assert screen.mutual[0].flow.classification == "mutual_suppression"
    with ZipFile(path) as archive:
        assert b"documentProtection" not in archive.read("word/settings.xml")


def test_ae1_embedded_chart_has_break_caption_and_resolvable_record_set(ae1):
    workspace, _, screen, path = ae1
    document = Document(path)
    assert len(document.inline_shapes) == 1
    caption = next(p.text for p in document.paragraphs if p.style.name == "Caption")
    assert "2019–2020" in caption
    assert "expected" in caption.lower() and "observed" in caption.lower()
    assert f"screens:{screen.screen_id}" in caption
    with ZipFile(path) as archive:
        images = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(images) == 1
        assert archive.read(images[0]).startswith(b"\x89PNG\r\n\x1a\n")
    parameters = json.loads(
        workspace.connection.execute(
            "SELECT parameters_json FROM screens WHERE id = ?",
            [screen.screen_id],
        ).fetchone()[0]
    )
    assert len(parameters["references"]) > 1  # Includes gravity-model peer inputs.
    refs = parameters["references"][0]
    for table, key in (
        ("positions", "position_ids"),
        ("workforce_inputs", "workforce_input_ids"),
        ("provenance", "provenance_ids"),
    ):
        assert refs[key]
        assert workspace.connection.execute(
            f"SELECT count(*) FROM {table} WHERE id IN (SELECT unnest(?))",
            [refs[key]],
        ).fetchone()[0] == len(refs[key])


def test_chart_plots_both_actual_directional_series(ae1, tmp_path, monkeypatch):
    workspace, _, screen, _ = ae1
    figures = []
    original = Figure.savefig

    def capture(self, *args, **kwargs):
        figures.append(self)
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Figure, "savefig", capture)
    export_dossier(workspace, screen, screen.mutual[0], tmp_path / "chart.docx")
    assert len(figures) == 1 and len(figures[0].axes) == 2
    for axis, direction in zip(
        figures[0].axes, screen.mutual[0].flow.directions, strict=True
    ):
        lines = {line.get_label(): line for line in axis.lines}
        assert set(lines) == {"Observed moves", "Expected moves"}
        assert list(lines["Observed moves"].get_xdata()) == list(range(2015, 2025))
        assert list(lines["Observed moves"].get_ydata()) == [
            y.observed for y in direction.years
        ]
        assert list(lines["Expected moves"].get_ydata()) == [
            y.expected for y in direction.years
        ]
        assert axis.patches  # Shaded break interval.


def test_no_clauses_is_explicit_and_partial_banner_is_visible(ae1):
    text = text_of(Document(ae1[3]))
    assert "no filed clauses found in the collected period" in text
    assert "Partial screen" in text
    assert "collection evidence" in text
    assert "synthetic" in text


def test_explicit_partial_parameter_displays_banner(ae1, tmp_path):
    from dataclasses import replace

    workspace, _, screen, _ = ae1
    complete = replace(screen, partial_screen=False)
    path = export_dossier(
        workspace,
        complete,
        complete.mutual[0],
        tmp_path / "partial.docx",
        partial_screen=True,
    )
    assert "Partial screen" in text_of(Document(path))


def test_filed_clause_citation_ties_and_declared_provenance(ae1, tmp_path):
    workspace, market, _, _ = ae1
    import_records(
        workspace,
        [
            {
                "person_key": "synthetic-director",
                "person_name": "Synthetic Director",
                "employer_name": name,
                "title": "Director",
                "location": "Fixture",
                "start_date": "2018-01-01",
                "end_date": "2024-12-31",
                "start_precision": "day",
                "end_precision": "day",
            }
            for name in ("Synthetic Amber Hospital", "Synthetic Birch Hospital")
        ],
        source="Synthetic consent fixture",
        provenance_class="consent_based",
    )
    employers = tuple(
        Employer(*r)
        for r in workspace.connection.execute(
            "SELECT id, name FROM employers ORDER BY id"
        ).fetchall()
    )
    receipt = workspace.provenance.append(
        source="synthetic",
        url="file:synthetic-clause",
        method="IMPORT",
        body=b"fixture clause",
        collection_method="synthetic",
    )
    clause_text = (
        "Synthetic Amber Hospital and Synthetic Birch Hospital: neither party "
        "shall hire the other party's employees for two years."
    )
    store_document(
        workspace,
        ClauseDocument(
            clause_text,
            receipt,
            date(2020, 1, 1),
            "synthetic_contract",
            employers[0],
            "Synthetic contract",
        ),
        employers,
    )
    screen = run_screen(workspace, market, accept_partial=True)
    pair = screen.mutual[0]
    path = export_dossier(workspace, screen, pair, tmp_path / "evidence.docx")
    text = text_of(Document(path))
    assert "neither party shall hire" in text
    assert "no filed clauses found" not in text
    assert "Synthetic contract" in text and "2020-01-01" in text
    assert "Synthetic Director" in text
    assert "consent_based" in text
    for clause in pair.clauses.clauses:
        assert f"clauses:{clause.id}" in text
        assert f"provenance:{clause.provenance.id}" in text
    for tie in pair.ties.ties:
        for position in tie.positions:
            assert f"positions:{position.id}" in text


def test_ae5_exported_text_passes_banned_phrase_scan(ae1):
    # Use the policy itself; matching any banned claim fails this acceptance gate.
    import runpy

    policy = runpy.run_path(str(Path(__file__).parents[1] / "rank" / "test_wording.py"))
    assert policy["BANNED"].search(text_of(Document(ae1[3]))) is None
