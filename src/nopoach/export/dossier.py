"""Editable Word pair dossiers from the persisted U10b screen result (U13).

Styles are built in code rather than stored as an opaque binary template.
No collection or signal recomputation occurs during export. A figure's stored
screen ID resolves the wider market inputs needed for its expected-flow model.
"""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure

from nopoach.market.sufficiency import Verdict
from nopoach.pipeline import ScreenResult
from nopoach.rank import wording
from nopoach.rank.aggregate import RankedPair
from nopoach.store.workspace import Workspace


def _document() -> WordDocument:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name, size in (
        ("Normal", 11),
        ("Title", 22),
        ("Heading 1", 14),
        ("Caption", 10),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.08
    document.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    document.core_properties.title = wording.DOSSIER_TITLE
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    return document


def _flow_chart(document: WordDocument, screen: ScreenResult, pair: RankedPair) -> None:
    if pair.flow is None or not any(d.years for d in pair.flow.directions):
        document.add_paragraph(wording.NO_FLOW)
        return
    document.add_paragraph(wording.FLOW_DESCRIPTION)
    names = {
        pair.employer_a_id: pair.employer_a_name,
        pair.employer_b_id: pair.employer_b_name,
    }
    figure = Figure(figsize=(6.8, 5.0), layout="constrained")
    FigureCanvasAgg(figure)
    axes = figure.subplots(2, 1, squeeze=False)
    breaks: list[tuple[str, tuple[int, int] | None]] = []
    for axis_row, direction in zip(axes, pair.flow.directions, strict=True):
        axis = axis_row[0]
        label = wording.flow_direction(
            names[direction.source_id], names[direction.destination_id]
        )
        years = [y.year for y in direction.years]
        axis.plot(
            years,
            [y.observed for y in direction.years],
            color="#111111",
            marker="o",
            markersize=3,
            label=wording.FLOW_OBSERVED,
        )
        axis.plot(
            years,
            [y.expected for y in direction.years],
            color="#385b80",
            linestyle="--",
            label=wording.FLOW_EXPECTED,
        )
        interval = (
            direction.break_result.year_interval if direction.break_result else None
        )
        if interval:
            axis.axvspan(*interval, color="#d4d4d4", alpha=0.65)
        axis.set_title(label, fontsize=10, loc="left")
        axis.set_ylabel(wording.FLOW_AXIS_MOVES, fontsize=9)
        axis.set_xlabel(wording.FLOW_AXIS_YEAR, fontsize=9)
        axis.set_xticks(years)
        axis.tick_params(labelsize=8)
        axis.set_ylim(bottom=0)
        axis.legend(loc="upper right", fontsize=8, frameon=False)
        axis.grid(axis="y", alpha=0.2)
        breaks.append((label, interval))
    image = BytesIO()
    figure.savefig(image, format="png", dpi=160)
    image.seek(0)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(image, width=Inches(6.8))
    document.add_paragraph(
        wording.flow_caption(
            wording.dossier_citation("screens", screen.screen_id), breaks
        ),
        style="Caption",
    )
    for label, direction in zip(
        (item[0] for item in breaks), pair.flow.directions, strict=True
    ):
        if direction.break_result:
            change = direction.break_result
            document.add_paragraph(
                wording.flow_break_description(
                    label,
                    change.before_observed_rate,
                    change.after_observed_rate,
                    change.year_interval,
                )
            )
    if not pair.flow.reliable:
        document.add_paragraph(pair.flow.reliability_reason)


def _coverage(document: WordDocument, screen: ScreenResult, pair: RankedPair) -> None:
    document.add_paragraph(wording.COVERAGE_DESCRIPTION)
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (2.1, 0.55, 1.2, 1.55, 1.4)
    for cell, label, width in zip(
        table.rows[0].cells, wording.COVERAGE_COLUMNS, widths, strict=True
    ):
        cell.text = label
        cell.width = Inches(width)
    for row in screen.coverage.employers:
        if row.employer.id not in (pair.employer_a_id, pair.employer_b_id):
            continue
        headcount = row.estimate.headcount if row.estimate else None
        values = (
            row.employer.name,
            str(row.year),
            str(row.visible_workers),
            f"{headcount.value:g}" if headcount else wording.SOURCE_CLASS_UNKNOWN,
            row.label,
        )
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths, strict=True):
            cell.text = value
            cell.width = Inches(width)
    # Light borders, wrapping rows and repeated headers keep annual data legible.
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            properties.append(borders)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E9E9E9")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = index == 0
    verdicts: dict[wording.SignalKey, Verdict] = {
        "flow": screen.sufficiency.flow,
        "clauses": screen.sufficiency.clauses,
        "ties": screen.sufficiency.ties,
    }
    for key, verdict in verdicts.items():
        document.add_paragraph(
            wording.sufficiency_entry(
                wording.SIGNAL_LABELS[key],
                verdict.status,
                verdict.reason,
            )
        )
    for reliability in screen.sufficiency.pairs:
        if (reliability.employer_a_id, reliability.employer_b_id) == (
            pair.employer_a_id,
            pair.employer_b_id,
        ):
            document.add_paragraph(reliability.reason)


def _sources(
    document: WordDocument, workspace: Workspace, screen: ScreenResult
) -> None:
    # Expected flow uses the wider market, so retain its receipts too.
    receipt_ids = {r for refs in screen.references for r in refs.provenance_ids}
    classes: dict[str, str] = {}
    if workspace.connection.execute(
        "SELECT table_name FROM information_schema.tables WHERE table_name = 'imports'"
    ).fetchone():
        classes = dict(
            workspace.connection.execute(
                "SELECT provenance_id, provenance_class FROM imports"
            ).fetchall()
        )
    for receipt in workspace.provenance.records():
        if receipt.id not in receipt_ids:
            continue
        classification = classes.get(receipt.id)
        if classification is None:
            classification = (
                "synthetic"
                if receipt.collection_method == "synthetic"
                else wording.SOURCE_CLASS_UNKNOWN
            )
        document.add_paragraph(
            wording.source_entry(
                receipt.source,
                classification,
                receipt.collection_method,
                receipt.url,
                receipt.collected_at.isoformat(),
                receipt.sha256,
                receipt.id,
            )
        )
        if classification == "synthetic":
            document.add_paragraph(wording.SYNTHETIC_NOTICE)
    document.add_paragraph(wording.RECORD_SET_NOTE)


def export_dossier(
    workspace: Workspace,
    screen: ScreenResult,
    pair: RankedPair,
    destination: str | Path,
    *,
    partial_screen: bool = False,
) -> Path:
    """Write a standalone dossier; U12 may explicitly request the partial banner.

    Pass a pair from this screen (including an unreliable group's pair), and
    the workspace in which run_screen persisted it. Explicit partial status
    adds the banner; the pipeline's partial status always remains visible.
    """
    document = _document()
    document.add_heading(wording.DOSSIER_TITLE, level=0)
    if screen.partial_screen or partial_screen:
        document.add_paragraph(wording.PARTIAL_SCREEN_BANNER).runs[0].bold = True
    document.add_heading(wording.DOSSIER_HEADINGS["summary"], level=1)
    document.add_paragraph(pair.summary)
    document.add_paragraph(
        wording.dossier_scope(
            pair.employer_a_name,
            pair.employer_b_name,
            sorted(screen.market.occupation_codes),
            screen.market.geography.kind,
            screen.market.geography.code,
            screen.coverage.market.start_year,
            screen.coverage.market.end_year,
            screen.as_of.isoformat() if screen.as_of else None,
        )
    )
    document.add_heading(wording.DOSSIER_HEADINGS["flow"], level=1)
    _flow_chart(document, screen, pair)
    document.add_heading(wording.DOSSIER_HEADINGS["clauses"], level=1)
    if pair.clauses and pair.clauses.clauses:
        for clause in pair.clauses.clauses:
            document.add_paragraph(clause.text)
            document.add_paragraph(
                wording.clause_citation(
                    clause.id,
                    clause.document_name,
                    clause.filed_on.isoformat(),
                    clause.start,
                    clause.end,
                    clause.binding,
                    clause.duration,
                    clause.scope,
                    clause.provenance.id,
                    clause.provenance.url,
                )
            )
    else:
        document.add_paragraph(wording.NO_FILED_CLAUSES)
    document.add_heading(wording.DOSSIER_HEADINGS["ties"], level=1)
    if pair.ties and pair.ties.ties:
        for tie in pair.ties.ties:
            document.add_paragraph(
                wording.tie_description(
                    tie.tie_type,
                    tie.start_date.isoformat() if tie.start_date else None,
                    tie.end_date.isoformat() if tie.end_date else None,
                )
            )
            for position in tie.positions:
                document.add_paragraph(
                    wording.tie_position(
                        position.person_name or position.person_id,
                        position.employer_name,
                        position.title,
                        f"{position.start_precision}/{position.end_precision}",
                        position.date_basis,
                        position.id,
                        position.provenance_id,
                    )
                )
    else:
        document.add_paragraph(wording.NO_TIES)
    document.add_heading(wording.DOSSIER_HEADINGS["explanations"], level=1)
    document.add_paragraph(screen.explanation_method)
    for check in pair.explanations:
        document.add_paragraph(
            wording.explanation_entry(
                wording.EXPLANATION_LABELS[check.key],
                wording.STATUS_LABELS[check.status],
                check.reason,
                check.provenance_ids,
            )
        )
    document.add_heading(wording.DOSSIER_HEADINGS["coverage"], level=1)
    _coverage(document, screen, pair)
    document.add_heading(wording.DOSSIER_HEADINGS["sources"], level=1)
    _sources(document, workspace, screen)
    document.add_heading(wording.DOSSIER_HEADINGS["method"], level=1)
    document.add_paragraph(wording.DOSSIER_METHOD)
    document.add_paragraph(screen.method)
    path = Path(destination)
    document.save(str(path))
    return path
